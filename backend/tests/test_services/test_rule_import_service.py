"""
规则导入服务纯函数测试（无DB依赖）
覆盖：行校验 / 章节切分 / 草稿标注 / 文档解析（docx表格、pdf文本层、OCR降级）
"""
import pytest

from app.config import settings
from app.services.rule_import_service import (
    ExtractedRule,
    annotate_draft_rows,
    read_document_text,
    split_sections,
    validate_rule_rows,
)


def _row(**over):
    """一行合法导入数据，覆盖部分字段造变体"""
    base = {
        "name": "测试规则",
        "code": "T1",
        "field_name": "amount",
        "operator": "gt",
        "threshold": "100",
    }
    base.update(over)
    return base


# ---------- validate_rule_rows ----------

def test_validate_rows_happy_path():
    """合法行：无错误（category_code→category_id 的解析在写入时由 to_rule_row 做）"""
    parsed, errs = validate_rule_rows(
        [_row(), _row(code="T2", category_code="meal")], set(), {"meal": 2}
    )
    assert errs == [[], []]
    assert parsed[0].category_id is None
    assert parsed[1].category_code == "meal"


def test_to_rule_row_resolves_category_code():
    """写入行构造：category_code 解析成 category_id，且不残留 category_code 键"""
    from app.services.rule_import_service import to_rule_row

    parsed, _ = validate_rule_rows([_row(code="T9", category_code="meal")], set(), {"meal": 2})
    row = to_rule_row(parsed[0], {"meal": 2})
    assert row["category_id"] == 2
    assert "category_code" not in row


def test_validate_rows_duplicate_with_db_and_batch():
    """code 与库内重复、批内重复都报错"""
    _, errs = validate_rule_rows([_row(code="DUP"), _row(code="DUP")], {"DUP"}, {})
    assert any("已存在" in e for e in errs[0])
    assert any("重复" in e for e in errs[1])


def test_validate_rows_invalid_category_code():
    """category_code 无效：错误信息附可用code列表"""
    _, errs = validate_rule_rows([_row(category_code="ghost")], set(), {"meal": 2})
    assert errs[0]
    assert any("meal" in e for e in errs[0])


def test_validate_rows_threshold_required_unless_exists():
    """非exists/not_exists操作符必须给阈值"""
    _, errs = validate_rule_rows([_row(threshold=None)], set(), {})
    assert any("threshold" in e or "阈值" in e for e in errs[0])

    _, errs2 = validate_rule_rows(
        [_row(field_name="invoice_no", operator="not_exists", threshold=None)], set(), {}
    )
    assert errs2[0] == []


def test_validate_rows_schema_violation_reported_per_row():
    """schema级违规（risk_points超界）进同一份行错误"""
    _, errs = validate_rule_rows([_row(risk_points=999)], set(), {})
    assert errs[0]
    assert any("risk_points" in e for e in errs[0])


def test_blank_category_code_normalized_to_none():
    """空串category_code（前端下拉清空）归一为None=全类别，不算错误"""
    parsed, errs = validate_rule_rows([_row(category_code="")], set(), {"meal": 2})
    assert errs == [[]]
    assert parsed[0].category_code is None


# ---------- split_sections ----------

def test_split_sections_by_headings():
    """标题行开新节"""
    text = "第一章 总则\n本制度适用于全体员工。\n第二条 住宿标准\n一线城市每晚上限600元。"
    secs = split_sections(text)
    assert [s["title"] for s in secs] == ["第一章 总则", "第二条 住宿标准"]
    assert "本制度适用于全体员工。" in secs[0]["content"]
    assert "600元" in secs[1]["content"]


def test_split_sections_no_heading_fallback():
    """无任何标题：单节兜底，标题用fallback_title"""
    text = "纯正文内容，没有任何标题行。\n第二行正文。"
    secs = split_sections(text, fallback_title="制度文档")
    assert len(secs) == 1
    assert secs[0]["title"] == "制度文档"
    assert "纯正文内容" in secs[0]["content"]


def test_split_sections_preserves_text_verbatim():
    """每一行原文都在节内逐字保留（保真：进向量库的不许被改写，标题行进content首行）"""
    lines = [
        "第一章 总则",
        "本制度适用于全体员工。",
        "（一）适用范围说明。",
        "第二条 住宿标准",
        "一线城市每晚上限600元。",
    ]
    secs = split_sections("\n".join(lines))
    body = "\n".join(s["content"] for s in secs)
    for line in lines:
        assert line in body


# ---------- annotate_draft_rows ----------

def test_annotate_draft_rows_flags_issues():
    """草稿轻校验：问题标到issues[]不抛错，quote原样带出"""
    drafts = [
        ExtractedRule(name="好规则", code="NEW_1", field_name="amount",
                      operator="gt", threshold="500", quote="原文A"),
        ExtractedRule(name="坏规则", code="OLD_9", category_code="ghost",
                      field_name="amount", operator="gt", threshold="1", quote="原文B"),
    ]
    out = annotate_draft_rows(drafts, {"OLD_9"}, {"meal": 2})
    assert out[0].issues == []
    assert out[0].quote == "原文A"
    assert any("OLD_9" in i for i in out[1].issues)
    assert any("meal" in i for i in out[1].issues)


# ---------- read_document_text ----------

def test_read_docx_full_includes_tables(tmp_path):
    """docx按文档顺序读段落+表格（表格单元格｜连接）"""
    from docx import Document

    doc = Document()
    doc.add_paragraph("第一章 报销标准")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "级别"
    table.cell(0, 1).text = "住宿上限"
    table.cell(1, 0).text = "一线城市"
    table.cell(1, 1).text = "600元"
    p = tmp_path / "t.docx"
    doc.save(str(p))

    text, method = read_document_text(p)
    assert method == "docx"
    assert "报销标准" in text
    assert "600元" in text and "一线城市" in text


def test_read_pdf_text_layer(tmp_path):
    """原生PDF走文本层直提（非OCR）"""
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    # 单行超页宽会被insert_text截断，分多行写；总长需超过MIN_PDF_TEXT_CHARS(200)
    # 才走文本层直提，否则被当扫描件降级OCR
    for i in range(4):
        page.insert_text(
            (72, 72 + 20 * i),
            f"Line {i}: hotel limit 600 yuan per night, detail clause {i} here.",
        )
    p = tmp_path / "t.pdf"
    doc.save(str(p))
    doc.close()

    text, method = read_document_text(p)
    assert method == "pdf_text"
    assert "600" in text


def test_read_pdf_fallback_to_ocr(tmp_path, monkeypatch):
    """文本层过短的PDF降级RapidOCR"""
    import pymupdf

    import app.services.rule_import_service as svc

    doc = pymupdf.open()
    doc.new_page()
    p = tmp_path / "blank.pdf"
    doc.save(str(p))
    doc.close()

    monkeypatch.setattr(svc, "run_ocr", lambda path: ("OCR识别出的制度文本" + "x" * 200, 0.9))
    text, method = read_document_text(p)
    assert method == "pdf_ocr"
    assert text.startswith("OCR识别出的制度文本")


# ---------- 抽取LLM客户端配置 ----------

def test_extraction_llm_timeout_bounded():
    """抽取LLM超时必须有界（Celery任务内调用也要兜底，无界挂起=worker被冻住）；
    用抽取专属的10分钟（思考模型大文档思考+工具调用超5分钟常见）；
    网络层不自动重试（重试把最坏等待翻倍），None截断由service层针对性重试"""
    from app.services.rule_import_service import (
        EXTRACTION_LLM_TIMEOUT_SECONDS,
        _build_extraction_llm,
    )

    llm = _build_extraction_llm()
    assert llm.request_timeout == EXTRACTION_LLM_TIMEOUT_SECONDS
    assert llm.max_retries == 0
    # 输出额度必须给足：思考token计入额度，8192会在思考阶段烧完导致工具调用被截断
    assert llm.max_tokens == 65536


def test_extract_rules_retries_once_on_none(monkeypatch):
    """思考模型烧完输出额度时不发工具调用（结构化输出None）：自动重试一次可救回"""
    from app.services import rule_import_service as svc

    calls = {"n": 0}

    class FakeStructured:
        def invoke(self, msgs):
            calls["n"] += 1
            if calls["n"] == 1:
                return None  # 第一次：思考token耗尽被截断
            return svc.ExtractedRules(
                rules=[
                    svc.ExtractedRule(
                        name="住宿上限", code="HOTEL_600", field_name="amount",
                        operator="gt", threshold="600",
                    )
                ]
            )

    class FakeLLM:
        def with_structured_output(self, schema, method=None, tool_choice=None):
            return FakeStructured()

    monkeypatch.setattr(svc, "_build_extraction_llm", lambda: FakeLLM())
    result = svc.extract_rules_from_text("制度文本", [])
    assert calls["n"] == 2
    assert result.rules[0].code == "HOTEL_600"


def test_extract_rules_none_twice_actionable_error(monkeypatch):
    """两次都None：报可操作的错误（提示重试/拆分文档），不再抛晦涩的"类型异常\""""
    import pytest

    from app.services import rule_import_service as svc

    class FakeStructured:
        def invoke(self, msgs):
            return None

    class FakeLLM:
        def with_structured_output(self, schema, method=None, tool_choice=None):
            return FakeStructured()

    monkeypatch.setattr(svc, "_build_extraction_llm", lambda: FakeLLM())
    with pytest.raises(svc.LLMExtractionError, match="拆分"):
        svc.extract_rules_from_text("制度文本", [])
